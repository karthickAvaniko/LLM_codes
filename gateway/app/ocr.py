"""
ocr.py — text extraction from any file type.

Images / scanned PDF pages go to the isolated OCR microservice (port 7780);
digital PDFs, Word, Excel and text are parsed directly.
"""
import io

import fitz  # PyMuPDF
import httpx
from PIL import Image

from app import core
log = core.log


def ocr_image(img_bytes: bytes, label: str = "") -> str:
    """Blocking call to the OCR service. Returns "" on any failure so the
    document pipeline degrades gracefully instead of erroring."""
    try:
        r = httpx.post(f"{core.OCR_URL}/ocr", content=img_bytes, timeout=120)
        if r.status_code == 200:
            text = r.json().get("text", "")
            log.info(f"OCR | {label} | {len(text)} chars")
            return text
        log.error(f"OCR service {r.status_code} | {label}")
    except Exception as e:
        log.error(f"OCR service unreachable | {label} | {e}")
    return ""


def extract_pages(content: bytes, filename: str) -> list[str]:
    """Extract text page-by-page. Pages with no embedded text get OCR'd
    individually, so mixed digital/scanned PDFs work."""
    ext = filename.lower().rsplit(".", 1)[-1]

    if ext in ("png", "jpg", "jpeg", "webp", "bmp", "tiff"):
        return [ocr_image(content, filename)]

    if ext == "pdf":
        doc   = fitz.open(stream=content, filetype="pdf")
        pages = []
        for i, page in enumerate(doc):
            text = page.get_text()
            if len(text.strip()) < 50:  # scanned / image page → OCR it
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
                buf = io.BytesIO()
                Image.frombytes("RGB", [pix.width, pix.height], pix.samples).save(buf, "PNG")
                text = ocr_image(buf.getvalue(), f"{filename} p{i+1}")
            pages.append(text)
        doc.close()
        return pages

    if ext == "docx":
        try:
            from docx import Document
            d = Document(io.BytesIO(content))
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            text = "\n".join(parts)
            return [text[i:i + 10_000] for i in range(0, len(text), 10_000)] or [""]
        except Exception as e:
            log.error(f"DOCX parse failed | {filename} | {e}")
            return [""]

    if ext in ("xlsx", "xls"):
        try:
            import pandas as pd
            sheets = pd.read_excel(io.BytesIO(content), sheet_name=None)
            return [f"[Sheet: {name}]\n{df.to_csv(index=False)}"
                    for name, df in sheets.items()] or [""]
        except Exception as e:
            log.error(f"XLSX parse failed | {filename} | {e}")
            return [""]

    # plain text / csv / etc. — split into ~10k-char pseudo-pages
    try:
        text = content.decode("utf-8", errors="ignore")
        return [text[i:i + 10_000] for i in range(0, len(text), 10_000)] or [""]
    except Exception:
        return [""]


def extract_text(content: bytes, filename: str) -> str:
    return "\n\n".join(extract_pages(content, filename))


def warm_ocr():
    """OCR lives in its own service (ocr_server.py) — nothing to warm here."""
    pass
